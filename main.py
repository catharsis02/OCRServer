import logging
import io
from pathlib import Path
import shutil
import uuid

from contextlib import asynccontextmanager
from fastapi import (
    FastAPI,
    UploadFile,
    File,
    HTTPException, Request,
)
from fastapi import status
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel

from docx import Document
from docx.shared import Pt

from backend.text_detection import TextDetection
from backend.text_recognition import TextRecognition
from backend.layout_detection import LayoutDetection



logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s"
)

logger = logging.getLogger(__name__)

@asynccontextmanager
async def lifespan(app: FastAPI):
    logger.info("Loading OCR models...")

    app.state.text_det = TextDetection(
        model_path=Path("./backend/openvino/text_detection/model.xml")
    )
    [shutil.rmtree(item) for item in ["det_outputs","pdf2image","temp_images"]]
    app.state.text_rec = TextRecognition(
        model_path=Path("./backend/openvino/text_recognition/model.xml"),
    )
    logger.info("Models loaded successfully")
    yield
    logger.info("Application shutting down")

app = FastAPI(
    title="OCR API",
    version="1.0.0",
    lifespan=lifespan
)

templates = Jinja2Templates(directory="./templates")
app.mount("/static",StaticFiles(directory="./static"),name="static")
app.mount("/fonts", StaticFiles(directory="./fonts"), name="fonts")

@app.get("/api/health")
def health_check():
    return {
        "status": "healthy"
    }

async def detect(
    file: UploadFile = File(...)
):

    allowed_types = {
        "application/pdf",
        "image/jpeg",
        "image/png"
    }

    if file.content_type not in allowed_types:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type: {file.content_type}"
        )

    job_id = str(uuid.uuid4())

    workspace = Path.cwd() / "jobs" / job_id

    uploads_dir = workspace / "uploads"
    pdf_to_photo_dir = workspace / "pdf_to_photo"
    det_save_dir = workspace / "det_save_path"
    rec_output_dir = workspace / "rec_outputs"
    final_text_dir = workspace / "final_text"
    temp_dir = workspace / "temp"

    try:
        uploads_dir.mkdir(
            parents=True,
            exist_ok=True
        )
        if file.filename:
            file_path = uploads_dir / Path(file.filename)

        logger.info(
            f"Job {job_id}: Processing {file.filename}"
        )

        content = await file.read()

        with open(file_path, "wb") as f:
            f.write(content)
        
        text_det = TextDetection(
            model_path=Path(
                "./backend/openvino/text_detection/model.xml"
            ),
            output_save_path=det_save_dir,
            temp_save_path=temp_dir,
            pdf_to_photo=pdf_to_photo_dir
        )

        text_rec = app.state.text_rec

        lay_dec = LayoutDetection(
            outputs_path=rec_output_dir / "json",
            out_text=final_text_dir,
            inflated_outputs_path=rec_output_dir / "inflated_boxes"
        )

        text_det.predict(
            pdf_path=uploads_dir
        )

        text_rec.predict(
            image_path=pdf_to_photo_dir,
            boxes_path=det_save_dir / "json",
            output_save_path=rec_output_dir
        )
        final_texts = lay_dec.predict()
        logger.info(
            f"Job {job_id}: OCR completed successfully"
        )
        return JSONResponse(
            content={
                "status": "success",
                "job_id": job_id,
                "total_lines": len(final_texts),
                "final_texts": final_texts
            }
        )

    except Exception as e:
        logger.exception(
            f"Job {job_id}: OCR pipeline failed"
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=str(e)
        )
    finally:
        logger.info(
            f"Job {job_id}: Cleaning workspace"
        )
        shutil.rmtree(
            workspace,
            ignore_errors=True
        )
    
@app.post("/api/detect",response_class=JSONResponse)
async def detect_api(
    file:UploadFile = File(...)
):
    return await detect(file)
    
@app.get("/", include_in_schema=False, name="home")
async def home(request:Request):
    return templates.TemplateResponse(request,"index.html")


class DocxRequest(BaseModel):
    text: str


def _build_docx(text: str) -> bytes:
    doc = Document()

    for line in text.split("\n"):
        p = doc.add_paragraph()

        if line.startswith("Page :"):
            run = p.add_run(line)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        else:
            run = p.add_run(line)
            run.font.name = "Calibri"
            run.font.size = Pt(12) 
            p.paragraph_format.space_after = Pt(3)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


@app.post("/api/download-docx")
async def download_docx(req: DocxRequest):
    if not req.text.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No text provided."
        )

    docx_bytes = _build_docx(req.text)

    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={
            "Content-Disposition": 'attachment; filename="ocr_output.docx"'
        },
    )